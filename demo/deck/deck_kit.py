"""What every slide of the pls,fix demo deck shares: the slide geometry the pane places into (it
must equal src/link/status.ts SLIDE_16_9 and src/ppt/placement.ts SLIDE_MARGIN/GAP), the Highlighter
colours, and the helpers that put a title, a frame, a box or a note on a slide. Python: no Rust
PPTX crate."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Pt

# Slide geometry in points: src/link/status.ts SLIDE_16_9, src/ppt/placement.ts
# SLIDE_MARGIN and SLIDE_GAP. The content area is the slide minus the margin on
# every side; halves split that area with the gap between them.
SLIDE_WIDTH = Pt(960)
SLIDE_HEIGHT = Pt(540)
MARGIN = Pt(36)
GAP = Pt(12)

CONTENT_LEFT = MARGIN
CONTENT_TOP = MARGIN
CONTENT_WIDTH = SLIDE_WIDTH - 2 * MARGIN
CONTENT_HEIGHT = SLIDE_HEIGHT - 2 * MARGIN

HALF_WIDTH = (CONTENT_WIDTH - GAP) // 2  # exact: divides evenly, never a float
HALF_HEIGHT = (CONTENT_HEIGHT - GAP) // 2
LEFT_HALF = (CONTENT_LEFT, CONTENT_TOP, HALF_WIDTH, CONTENT_HEIGHT)
RIGHT_HALF = (CONTENT_LEFT + HALF_WIDTH + GAP, CONTENT_TOP, HALF_WIDTH, CONTENT_HEIGHT)
# The four quarters (src/layout.ts spotBox: the same halves on both axes).
TOP_LEFT = (CONTENT_LEFT, CONTENT_TOP, HALF_WIDTH, HALF_HEIGHT)
TOP_RIGHT = (CONTENT_LEFT + HALF_WIDTH + GAP, CONTENT_TOP, HALF_WIDTH, HALF_HEIGHT)
BOTTOM_LEFT = (CONTENT_LEFT, CONTENT_TOP + HALF_HEIGHT + GAP, HALF_WIDTH, HALF_HEIGHT)
BOTTOM_RIGHT = (CONTENT_LEFT + HALF_WIDTH + GAP, CONTENT_TOP + HALF_HEIGHT + GAP, HALF_WIDTH, HALF_HEIGHT)

# The pane's free-space scan (src/ppt/placement.ts occupiedBoxes, src/layout.ts
# placeInFreeSpace) blocks anything whose top is less than a shape's bottom + GAP,
# so the title (text + rule) must end by here for CONTENT_TOP to read as free.
TITLE_BAND_LIMIT = CONTENT_TOP - GAP  # 24 pt
RULE_HEIGHT = Pt(2)
TITLE_HEIGHT = TITLE_BAND_LIMIT - RULE_HEIGHT  # 22 pt: right up to the rule

# Slide 4's caption band: CAPTION_HEIGHT is the one named constant, GAP is
# the breathing room above it, so picture + GAP + caption ties out to
# CONTENT_HEIGHT exactly instead of three independently hand-tuned numbers.
CAPTION_HEIGHT = Pt(40)
PICTURE_HEIGHT = CONTENT_HEIGHT - GAP - CAPTION_HEIGHT
CAPTION_TOP = CONTENT_TOP + CONTENT_HEIGHT - CAPTION_HEIGHT

# Brand colors, matching the panes (tokens.css palette, the Highlighter look). Ink
# draws every word, rule and outline here: lime (#D9F24A) is a fill-only colour (1.2:1
# on white, so never a line or text), and the dashed half-guides stay unfilled because
# the pane reads a no-fill dashed frame as free space (src/layout.ts isDecorativeFrame).
INK = RGBColor(0x11, 0x11, 0x11)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# The object-tool slides' fills: the theme's own accents (see THEME_COLORS in build_deck.py).
OLIVE = RGBColor(0x6D, 0x7E, 0x12)
GREY = RGBColor(0xCB, 0xCB, 0xCB)
TINT = RGBColor(0xD3, 0xD8, 0xB8)

TITLE_AND_CONTENT = 1
TITLE_ONLY = 5
PICTURE_WITH_CAPTION = 8


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return prs


def set_box(shape, left, top, width, height) -> None:
    shape.left, shape.top, shape.width, shape.height = left, top, width, height


def add_title(slide, text: str) -> None:
    """Put the title in the margin band (0..TITLE_BAND_LIMIT) above the
    content area and rule it off in ink, so the pane's free-space scan
    reads CONTENT_TOP as free rather than blocked by the title's own gap
    buffer. Vertical anchor and insets set explicitly - "Picture with
    Caption" bottom-anchors its title with a bottom inset, which clipped a
    16 pt line at the top of an 18 pt box - so no layout's own anchoring is
    inherited on any slide. Centered on every slide too: a layout's own
    alignment (Picture with Caption defaults to left) is not trusted."""
    title = slide.shapes.title
    title.text = text
    set_box(title, CONTENT_LEFT, Pt(0), CONTENT_WIDTH, TITLE_HEIGHT)
    tf = title.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_top = 0
    tf.margin_bottom = 0
    paragraph = tf.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    font = paragraph.font
    font.size = Pt(16)
    font.bold = True
    font.color.rgb = INK

    rule_top = TITLE_BAND_LIMIT - RULE_HEIGHT
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, CONTENT_LEFT, rule_top, CONTENT_WIDTH, RULE_HEIGHT)
    rule.fill.solid()
    rule.fill.fore_color.rgb = INK
    rule.line.fill.background()
    rule.shadow.inherit = False


def add_slide(prs: Presentation, layout: int, title_text: str):
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    add_title(slide, title_text)
    return slide


def set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def set_number_bullet(paragraph) -> None:
    """Real auto-numbering (1., 2., ...) instead of the master's inherited
    bullet char, reached via the paragraph's own oxml element - python-pptx
    has no public API for bullet formatting."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.remove_all("a:buNone", "a:buAutoNum", "a:buChar")
    pPr.insert_element_before(
        parse_xml(f'<a:buFont {nsdecls("a")} typeface="+mj-lt"/>'),
        "a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst",
    )
    pPr.insert_element_before(
        parse_xml(f'<a:buAutoNum {nsdecls("a")} type="arabicPeriod"/>'),
        "a:buNone", "a:buChar", "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst",
    )


def set_no_bullet(paragraph) -> None:
    """Suppress the master's inherited bullet on a plain paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.remove_all("a:buNone", "a:buAutoNum", "a:buChar")
    pPr.insert_element_before(
        parse_xml(f'<a:buNone {nsdecls("a")}/>'),
        "a:buAutoNum", "a:buChar", "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst",
    )


def add_frame(slide, box):
    """A dashed ink frame with no fill and no text: the pane reads it as free space, so an
    object inserted into its spot lands on it without counting it as an obstacle."""
    left, top, width, height = box
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    rect.fill.background()
    rect.line.color.rgb = INK
    rect.line.width = Pt(2)
    rect.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    rect.shadow.inherit = False
    return rect


def add_caption(slide, left, top, width, text: str, size: int = 12) -> None:
    """One centred line of ink text under (or beside) the exercise."""
    caption = slide.shapes.add_textbox(left, top, width, Pt(24))
    para = caption.text_frame.paragraphs[0]
    para.text = text
    para.alignment = PP_ALIGN.CENTER
    para.font.size = Pt(size)
    para.font.color.rgb = INK


def add_box(slide, name: str, box, label: str, fill: RGBColor, text: RGBColor = INK,
            shape=MSO_SHAPE.RECTANGLE, outline_pt: float = 1.25):
    """A named, labelled object for the object tools to act on: solid fill, ink outline, centred
    bold label, no theme shadow. The name is what the selection pane shows."""
    left, top, width, height = box
    obj = slide.shapes.add_shape(shape, left, top, width, height)
    obj.name = name
    obj.fill.solid()
    obj.fill.fore_color.rgb = fill
    obj.line.color.rgb = INK
    obj.line.width = Pt(outline_pt)
    obj.shadow.inherit = False
    para = obj.text_frame.paragraphs[0]
    para.text = label
    para.alignment = PP_ALIGN.CENTER
    para.font.size = Pt(16)
    para.font.bold = True
    para.font.color.rgb = text
    return obj
