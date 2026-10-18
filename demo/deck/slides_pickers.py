"""Slides 1-5 of the pls,fix demo deck: the overview and one landing spot for every link kind the
Excel Links tab sends. Slide 2 is a table's whole slide, slide 3 a chart's two halves, slide 4 a
picture's placeholder, slide 5 a text's quarter (and a table's and a chart's). The shape order on
slides 2-4 is fixed on purpose: the video capture deletes slide 3's guides by their automatic names."""

from pptx import Presentation
from pptx.util import Pt

from deck_kit import (
    BOTTOM_LEFT, BOTTOM_RIGHT, CAPTION_HEIGHT, CAPTION_TOP, CONTENT_HEIGHT, CONTENT_LEFT, CONTENT_TOP,
    CONTENT_WIDTH, INK, LEFT_HALF, PICTURE_HEIGHT, PICTURE_WITH_CAPTION, RIGHT_HALF, TITLE_AND_CONTENT,
    TITLE_ONLY, TOP_LEFT, TOP_RIGHT, add_caption, add_frame, add_slide, set_box, set_no_bullet, set_notes,
    set_number_bullet,
)

OVERVIEW_TASKS = [
    "P&L table (a table link) → Slide 2, Where = Whole slide",
    "Revenue chart (a chart link) → Slide 3, Where = Left half",
    "Segment pie (a chart link) → Slide 3, Where = Right half",
    "Picture from the Data sheet (a picture link) → Slide 4: select the empty "
    "placeholder, Where = Selected shape",
    "Headline text (a text link) → Slide 5, Where = Top left; a second text, the table and "
    "the chart fill the other quarters",
    "Object tools on shapes already drawn: Slide 6 align and distribute, Slide 7 match size, "
    "select similar and swap, Slide 8 Smart Painter",
]


def build_slide_1(prs: Presentation) -> None:
    """The tasks as a numbered list, one per link kind and tool group, plus how to run them."""
    slide = add_slide(prs, TITLE_AND_CONTENT, "pls,fix demo deck")
    body = slide.placeholders[1]
    set_box(body, CONTENT_LEFT, CONTENT_TOP, CONTENT_WIDTH, CONTENT_HEIGHT)
    tf = body.text_frame
    tf.word_wrap = True

    tf.paragraphs[0].text = OVERVIEW_TASKS[0]
    tf.paragraphs[0].font.size = Pt(18)
    set_number_bullet(tf.paragraphs[0])
    for task in OVERVIEW_TASKS[1:]:
        para = tf.add_paragraph()
        para.text = task
        para.font.size = Pt(18)
        set_number_bullet(para)

    commands = tf.add_paragraph()
    commands.space_before = Pt(20)
    set_no_bullet(commands)
    label = commands.add_run()
    label.text, label.font.bold, label.font.size = "Commands: ", True, Pt(14)
    rest = commands.add_run()
    rest.text = (
        "open the pane from the pls,fix ribbon tab, Inbox tab (Paste from Excel, the Slide and "
        "Where pickers above the list, Insert), Links tab (Update all), Tools tab for the shapes."
    )
    rest.font.size = Pt(14)
    set_notes(slide, "This deck demos the PowerPoint Slide and Where pickers across the four link kinds "
              "(table, chart, picture, text) and the object tools. Follow the Start here sheet of the "
              "demo workbook.")


def build_slide_2(prs: Presentation) -> None:
    """One empty content placeholder: free space for the pane to place into."""
    slide = add_slide(prs, TITLE_AND_CONTENT, "Choose the slide")
    body = slide.placeholders[1]
    set_box(body, CONTENT_LEFT, CONTENT_TOP, CONTENT_WIDTH, CONTENT_HEIGHT)
    # No text written: an empty layout placeholder reads as free space, the
    # same rule src/ppt/placement.ts applies when it scans a slide's shapes.
    set_notes(slide, "Choose the slide: pick this slide, then a spot, in the Slide picker.")


def add_half_diagram(slide, box, caption_text: str) -> None:
    left, top, width, height = box
    add_frame(slide, box)
    add_caption(slide, left, top + height + Pt(4), width, caption_text)


def build_slide_3(prs: Presentation) -> None:
    """The pane's left and right halves, drawn to their exact geometry."""
    slide = add_slide(prs, TITLE_ONLY, "Choose the spot")
    add_half_diagram(slide, LEFT_HALF, "Revenue chart: Where = Left half")
    add_half_diagram(slide, RIGHT_HALF, "Segment pie: Where = Right half")
    set_notes(slide, "Choose the spot: Where picks a half, a quarter, the whole slide or the selected shape.")


def build_slide_4(prs: Presentation) -> None:
    """An empty picture placeholder to select and insert into."""
    slide = add_slide(prs, PICTURE_WITH_CAPTION, "Into a placeholder")
    picture = slide.placeholders[1]
    set_box(picture, CONTENT_LEFT, CONTENT_TOP, CONTENT_WIDTH, PICTURE_HEIGHT)
    # Left empty: no insert_picture() call, so it stays the "click to add
    # picture" placeholder the pane's occupied-box scan treats as free space.

    caption = slide.placeholders[2]
    set_box(caption, CONTENT_LEFT, CAPTION_TOP, CONTENT_WIDTH, CAPTION_HEIGHT)
    para = caption.text_frame.paragraphs[0]
    para.text = "Select the placeholder, Where = Selected shape, Insert."
    para.font.size = Pt(14)
    para.font.color.rgb = INK
    set_notes(slide, "Into a placeholder: select the empty picture placeholder, Where = Selected shape, Insert.")


def build_slide_5(prs: Presentation) -> None:
    """The four quarters, drawn to their exact geometry, with one caption line under the grid."""
    slide = add_slide(prs, TITLE_ONLY, "Choose the spot: four quarters")
    for box in (TOP_LEFT, TOP_RIGHT, BOTTOM_LEFT, BOTTOM_RIGHT):
        add_frame(slide, box)
    add_caption(
        slide, CONTENT_LEFT, CONTENT_TOP + CONTENT_HEIGHT + Pt(4), CONTENT_WIDTH,
        "Top left and top right: the two text links; bottom left: the P&L table; bottom right: "
        "the Revenue chart",
    )
    set_notes(slide, "Choose the spot: a quarter is a Where option too. Insert each Excel export "
              "with its own quarter: the text boxes keep their place when they refresh.")
