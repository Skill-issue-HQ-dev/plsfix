"""Slides 6-8 of the pls,fix demo deck: shapes drawn and named for the PowerPoint object tools to act
on. Slide 6 is deliberately untidy (align, distribute), slide 7 holds equal and unequal boxes plus two
circles (select similar, match size, swap), slide 8 one styled box and three plain ones (Smart
Painter). Every shape is a named, labelled box so the selection pane reads A, B, C and so on."""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from deck_kit import (
    CONTENT_LEFT, CONTENT_WIDTH, GREY, OLIVE, TINT, TITLE_ONLY, WHITE, add_box, add_caption, add_slide, set_notes,
)

CAPTION_TOP_PT = Pt(430)
BOX = (Pt(110), Pt(60))


def box_at(left: float, top: float, size=BOX):
    """A (left, top, width, height) tuple from points."""
    return (Pt(left), Pt(top), size[0], size[1])


def build_slide_6(prs: Presentation) -> None:
    """Five equal boxes at ragged heights and uneven gaps: Align Top and Distribute Across have work to do."""
    slide = add_slide(prs, TITLE_ONLY, "Object tools: align and distribute")
    spots = {"A": (60, 150), "B": (180, 250), "C": (300, 130), "D": (560, 260), "E": (790, 170)}
    for name, (left, top) in spots.items():
        add_box(slide, f"Box {name}", box_at(left, top), name, GREY)
    add_caption(
        slide, CONTENT_LEFT, CAPTION_TOP_PT, CONTENT_WIDTH,
        "Select all five, choose Top in Align and press Apply. Then select A, C and E, choose Across "
        "in Distribute and press Apply.", size=14,
    )
    set_notes(slide, "Align and distribute: Align uses the selection's own bounds, Distribute needs "
              "three or more objects and keeps the outer two where they are. The ribbon's Arrange group "
              "does the same.")


def build_slide_7(prs: Presentation) -> None:
    """Three equal boxes and two odd ones, two circles: Select similar, Match size and Swap."""
    slide = add_slide(prs, TITLE_ONLY, "Object tools: match size, select similar, swap")
    # Match size gives A and E the width and height of D (200 x 110 pt). A grows to the right and E
    # down and to the right, so B starts 20 pt past A's new edge and E sits clear of D, both circles
    # and the slide's edge: nothing overlaps before or after the match.
    for name, left in (("A", 60), ("B", 280), ("C", 420)):
        add_box(slide, f"Box {name}", box_at(left, 110, (Pt(120), Pt(70))), name, GREY)
    add_box(slide, "Box D", box_at(580, 100, (Pt(200), Pt(110))), "D", TINT)
    add_box(slide, "Box E", box_at(440, 245, (Pt(90), Pt(50))), "E", TINT)
    circle = (Pt(90), Pt(90))
    add_box(slide, "Circle 1", box_at(180, 300, circle), "1", GREY, shape=MSO_SHAPE.OVAL)
    add_box(slide, "Circle 2", box_at(700, 300, circle), "2", TINT, shape=MSO_SHAPE.OVAL)
    add_caption(
        slide, CONTENT_LEFT, CAPTION_TOP_PT, CONTENT_WIDTH,
        "Select similar: click A. Match size: select D first, then A and E. Swap: select the two circles.",
        size=14,
    )
    set_notes(slide, "Select similar picks every object of the same kind and size as the one you clicked. "
              "Match size gives every target the width and height of the FIRST object you selected. "
              "Swap exchanges the places of exactly two objects.")


def build_slide_8(prs: Presentation) -> None:
    """One olive box with a heavy outline and three plain grey ones: capture, then apply."""
    slide = add_slide(prs, TITLE_ONLY, "Object tools: Smart Painter")
    add_box(slide, "Source", box_at(80, 140, (Pt(200), Pt(110))), "Source", OLIVE, text=WHITE, outline_pt=3)
    for index, left in enumerate((340, 520, 700), start=1):
        add_box(slide, f"Target {index}", box_at(left, 140, (Pt(150), Pt(110))), f"Target {index}", GREY)
    add_caption(
        slide, CONTENT_LEFT, CAPTION_TOP_PT, CONTENT_WIDTH,
        "Select the olive box and press Capture. Select the three grey boxes and press Apply: they take "
        "its fill, outline and text look, not its size or words.", size=14,
    )
    set_notes(slide, "Smart Painter copies a solid fill, an outline and the text look (font, size, color, "
              "bold, italic), never an object's content or geometry. Capture one object, then Apply to "
              "every selected target.")
